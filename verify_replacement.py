import json
from docx import Document

def verify():
    mapping_path = r'F:\AI\тест2\DB\Книги ачарий\verse_mapping.json'
    input_docx_path = r'F:\AI\тест2\DB\Книги ачарий\Бхакти-расамрита-синдху часть 1 (OCR).docx'
    output_docx_path = r'F:\AI\тест2\DB\Книги ачарий\Бхакти-расамрита-синдху часть 1 (OCR)_2.docx'

    with open(mapping_path, 'r', encoding='utf-8') as f:
        mapping = json.load(f)

    doc_old = Document(input_docx_path)
    doc_new = Document(output_docx_path)

    old_texts = [p.text.strip() for p in doc_old.paragraphs if p.text.strip()]
    new_texts = [p.text.strip() for p in doc_new.paragraphs if p.text.strip()]

    replaced_found = 0
    for old_text, new_text in mapping.items():
        if old_text in old_texts and new_text in new_texts:
            replaced_found += 1
    
    print(f"Found {replaced_found} replacements in the new document that match the mapping.")

if __name__ == "__main__":
    verify()
